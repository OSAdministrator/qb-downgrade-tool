"""Generate Excel + PDF report of memorized transactions for the operator.

QBFC SDK has no Add for memorized transactions — they're templates bound to
existing txns. We capture the list during export and produce two artifacts
(.xlsx and .pdf) so the operator can re-memorize them manually in QB 2021.
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple


def _columns() -> List[Tuple[str, str, int]]:
    """(json_key, display, excel_width)."""
    return [
        ("name",            "Name",             36),
        ("txn_type",        "Type",             16),
        ("how_often",       "Frequency",        14),
        ("next_date",       "Next Date",        14),
        ("days_in_advance", "Days In Advance",  16),
        ("remaining_times", "Remaining",        12),
        ("is_active",       "Active",           10),
    ]


def _fmt(val: Any) -> str:
    if val is None or val == "":
        return ""
    if isinstance(val, bool):
        return "Yes" if val else "No"
    return str(val)


def write_excel(memorized: List[Dict[str, Any]], path: Path, company_name: str) -> Path:
    """Write memorized transactions to an .xlsx file."""
    from openpyxl import Workbook  # type: ignore
    from openpyxl.styles import Alignment, Font, PatternFill  # type: ignore

    cols = _columns()

    wb = Workbook()
    ws = wb.active
    ws.title = "Memorized Transactions"

    # Title row
    ws.cell(row=1, column=1, value=f"Memorized Transactions \u2014 {company_name}")
    ws.cell(row=1, column=1).font = Font(size=14, bold=True, color="FFFFFF")
    ws.cell(row=1, column=1).fill = PatternFill("solid", fgColor="1E293B")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(cols))
    ws.cell(row=1, column=1).alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26

    # Subtitle row
    sub = (
        f"Generated {datetime.now().strftime('%Y-%m-%d %H:%M')}  \u2022  "
        f"{len(memorized)} item(s)  \u2022  Re-memorize in QB 2021 via Edit \u2192 Memorize"
    )
    ws.cell(row=2, column=1, value=sub)
    ws.cell(row=2, column=1).font = Font(size=9, italic=True, color="475569")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(cols))
    ws.cell(row=2, column=1).alignment = Alignment(horizontal="center")

    # Header row
    header_row = 4
    for idx, (_key, label, width) in enumerate(cols, start=1):
        c = ws.cell(row=header_row, column=idx, value=label)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="334155")
        c.alignment = Alignment(horizontal="left", vertical="center")
        ws.column_dimensions[c.column_letter].width = width
    ws.row_dimensions[header_row].height = 20

    # Data rows
    for r_idx, m in enumerate(memorized, start=header_row + 1):
        for c_idx, (key, _label, _w) in enumerate(cols, start=1):
            ws.cell(row=r_idx, column=c_idx, value=_fmt(m.get(key)))
        # Alternating row banding
        if r_idx % 2 == 0:
            for c_idx in range(1, len(cols) + 1):
                ws.cell(row=r_idx, column=c_idx).fill = PatternFill("solid", fgColor="F1F5F9")

    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))
    return path


def write_pdf(memorized: List[Dict[str, Any]], path: Path, company_name: str) -> Path:
    """Write memorized transactions to a printable .pdf file."""
    from reportlab.lib import colors  # type: ignore
    from reportlab.lib.pagesizes import landscape, letter  # type: ignore
    from reportlab.lib.styles import getSampleStyleSheet  # type: ignore
    from reportlab.lib.units import inch  # type: ignore
    from reportlab.platypus import (  # type: ignore
        Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
    )

    cols = _columns()
    styles = getSampleStyleSheet()

    path.parent.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(
        str(path),
        pagesize=landscape(letter),
        leftMargin=0.4 * inch, rightMargin=0.4 * inch,
        topMargin=0.4 * inch, bottomMargin=0.4 * inch,
        title=f"Memorized Transactions \u2014 {company_name}",
    )

    story = []
    story.append(Paragraph(
        f"<b>Memorized Transactions \u2014 {company_name}</b>", styles["Title"]))
    story.append(Paragraph(
        f"Generated {datetime.now().strftime('%Y-%m-%d %H:%M')} "
        f"\u2022 {len(memorized)} item(s) "
        f"\u2022 Re-memorize in QB 2021 via <b>Edit \u2192 Memorize</b> on each matching transaction.",
        styles["Italic"]))
    story.append(Spacer(1, 12))

    headers = [c[1] for c in cols]
    data = [headers]
    for m in memorized:
        data.append([_fmt(m.get(c[0])) for c in cols])

    if len(data) == 1:
        data.append(["(no memorized transactions found)"] + [""] * (len(cols) - 1))

    tbl = Table(data, repeatRows=1)
    tbl.setStyle(TableStyle([
        ("BACKGROUND",   (0, 0), (-1, 0), colors.HexColor("#334155")),
        ("TEXTCOLOR",    (0, 0), (-1, 0), colors.white),
        ("FONTNAME",     (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE",     (0, 0), (-1, 0), 10),
        ("BOTTOMPADDING",(0, 0), (-1, 0), 6),
        ("FONTSIZE",     (0, 1), (-1, -1), 9),
        ("VALIGN",       (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN",        (0, 0), (-1, -1), "LEFT"),
        ("ROWBACKGROUNDS",(0, 1),(-1, -1), [colors.white, colors.HexColor("#F1F5F9")]),
        ("GRID",         (0, 0), (-1, -1), 0.25, colors.HexColor("#CBD5E1")),
        ("LEFTPADDING",  (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
    ]))
    story.append(tbl)

    doc.build(story)
    return path


def generate_reports(
    memorized: List[Dict[str, Any]],
    output_dir: Path,
    company_name: str,
) -> Dict[str, Optional[Path]]:
    """Generate both Excel and PDF reports. Returns dict with paths."""
    output_dir.mkdir(parents=True, exist_ok=True)
    safe = "".join(ch for ch in company_name if ch.isalnum() or ch in " -_").strip() or "company"
    xlsx_path = output_dir / f"{safe} - Memorized Transactions.xlsx"
    pdf_path  = output_dir / f"{safe} - Memorized Transactions.pdf"

    out: Dict[str, Optional[Path]] = {"xlsx": None, "pdf": None}
    try:
        out["xlsx"] = write_excel(memorized, xlsx_path, company_name)
    except Exception:
        out["xlsx"] = None
    try:
        out["pdf"] = write_pdf(memorized, pdf_path, company_name)
    except Exception:
        out["pdf"] = None
    return out


__all__ = ["generate_reports", "write_excel", "write_pdf"]
